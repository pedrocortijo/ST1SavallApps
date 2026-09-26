from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path(__file__).with_name('requisitos_aplicacion_hardware_software.docx')
BLUE, GRID = '17365D', 'D9D9D9'

def run_style(run, bold=False, size=10.2, color=None):
    run.font.name = 'Aptos'; run.font.size = Pt(size); run.bold = bold
    run._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos'); run._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
    if color: run.font.color.rgb = RGBColor.from_string(color)

def cell_fmt(cell, fill=None):
    tcpr = cell._tc.get_or_add_tcPr()
    if fill:
        shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tcpr.append(shd)
    borders = OxmlElement('w:tcBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        e = OxmlElement(f'w:{edge}'); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'6'); e.set(qn('w:color'),GRID); borders.append(e)
    tcpr.append(borders)
    mar = OxmlElement('w:tcMar')
    for side in ('top','start','bottom','end'):
        e = OxmlElement(f'w:{side}'); e.set(qn('w:w'),'115'); e.set(qn('w:type'),'dxa'); mar.append(e)
    tcpr.append(mar); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def p(doc, text, bullet=False):
    x = doc.add_paragraph(style='List Bullet' if bullet else None)
    x.paragraph_format.space_after = Pt(4); x.paragraph_format.line_spacing = 1.1
    run_style(x.add_run(text)); return x

def h(doc, text):
    x = doc.add_paragraph(style='Heading 1'); x.paragraph_format.space_before = Pt(12); x.paragraph_format.space_after = Pt(5)
    run_style(x.add_run(text), True, 13.5, '000000'); return x

def table(doc, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers)); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i, label in enumerate(headers):
        c=t.rows[0].cells[i]; c.text=''; cell_fmt(c, BLUE); run_style(c.paragraphs[0].add_run(label), True, 9.2, 'FFFFFF')
    for r, row in enumerate(rows):
        cells=t.add_row().cells
        for i, value in enumerate(row):
            c=cells[i]; c.text=''; cell_fmt(c, 'F4F7FA' if r%2 else None); q=c.paragraphs[0]; q.paragraph_format.space_after=Pt(0); run_style(q.add_run(value), False, 9.1)
    for row in t.rows:
        for i, width in enumerate(widths): row.cells[i].width=Cm(width)
    doc.add_paragraph().paragraph_format.space_after=Pt(2)

d=Document(); s=d.sections[0]; s.top_margin=Cm(2.1); s.bottom_margin=Cm(2.0); s.left_margin=Cm(2.15); s.right_margin=Cm(2.15)
d.styles['Normal'].font.name='Aptos'; d.styles['Normal'].font.size=Pt(10.2)
title=d.add_paragraph(style='Title'); title.paragraph_format.space_after=Pt(8); run_style(title.add_run('Requisitos clave de la aplicación hardware y software'), True, 21, '000000')
sub=d.add_paragraph(); sub.paragraph_format.space_after=Pt(16); run_style(sub.add_run('Aplicación móvil ST1 Savall | Versión inicial | 19 de septiembre de 2026'), False, 11, '595959')
p(d,'Este documento define el equipo mínimo, el sistema operativo, el software de soporte y las condiciones de operación para desplegar la aplicación móvil ST1 Savall en personal de campo. Los requisitos se orientan a una operación fiable de servicios, captura de datos, firma, fotografías y comunicación con la API corporativa.')

h(d,'1 Plataforma compatible')
table(d,['Elemento','Requisito','Condición de aceptación'],[
['Sistema operativo','Android 7.0 o superior, API level 24.','La aplicación se instala, inicia y completa una prueba funcional en el dispositivo.'],
['Arquitectura','Dispositivo Android certificado por el fabricante y con servicios de Google si la distribución los requiere.','Sin bloqueo del sistema, root ni ROM no gestionada.'],
['Formato','Tableta rugerizada o teléfono corporativo según el puesto de trabajo.','La pantalla, cámara y uso con funda se ajustan al trabajo real.'],
['Gestión','Alta en MDM o inventario corporativo antes de entregar el equipo.','Equipo asignado a usuario, área y responsable de soporte.']], [3.0,7.2,5.2])

h(d,'2 Requisitos de hardware')
table(d,['Componente','Mínimo','Recomendado para operación de campo'],[
['Procesador y memoria','4 núcleos y 3 GB de RAM.','8 núcleos y 4 GB de RAM o superior.'],
['Almacenamiento','32 GB internos, con al menos 8 GB libres al entregar.','64 GB o más, para fotografías, documentos y actualizaciones.'],
['Pantalla','5,5 pulgadas, resolución HD y brillo legible en interior.','8 pulgadas o más, Full HD, alto brillo y respuesta táctil con guantes.'],
['Cámara','Cámara trasera funcional con enfoque automático.','12 MP o superior, flash y lectura fiable de documentos o códigos.'],
['Conectividad','Wi Fi y 4G LTE. Bluetooth si se usan accesorios.','4G o 5G, Wi Fi de doble banda, Bluetooth 5 y GPS asistido.'],
['Ubicación','GPS operativo y permiso de ubicación disponible.','GNSS con buena recepción y actualización rápida en exterior.'],
['Batería','Autonomía de una jornada o carga disponible en vehículo.','Batería de 6.000 mAh o superior y cargador de vehículo homologado.'],
['Resistencia','Funda protectora y protección frente a golpes y polvo.','IP65 o superior, protección contra caída y uso entre -10 y 45 grados C.']], [3.1,5.8,6.5])

h(d,'3 Software base y configuración')
p(d,'Los dispositivos se entregarán con una versión estable y soportada de Android, actualizaciones de seguridad disponibles y únicamente las aplicaciones necesarias para la función asignada.')
table(d,['Área','Requisito'],[
['Aplicación','Instalar la versión aprobada de ST1 Savall desde el canal de distribución corporativo y registrar versión, fecha y dispositivo.'],
['Servicios Android','Habilitar cámara, ubicación, almacenamiento y notificaciones solo cuando el proceso de negocio los necesite.'],
['Navegador y documentos','Mantener un navegador actualizado y visor PDF para consultar documentos generados por la aplicación.'],
['Correo y mensajería','Usar solo clientes corporativos aprobados; no reenviar documentos de trabajo a cuentas personales.'],
['Gestión remota','Aplicar bloqueo de pantalla, cifrado, inventario, borrado remoto y actualización gestionada mediante MDM cuando esté disponible.'],
['Antivirus','Usar protección móvil corporativa si la política de seguridad lo exige o si el dispositivo permite instalar software fuera del canal aprobado.']], [3.6,11.8])

h(d,'4 Red y conectividad')
p(d,'La aplicación necesita acceso HTTPS a la API corporativa para las operaciones en línea. Antes de implantar, se debe validar la URL productiva desde redes Wi Fi y móviles habituales, así como el tratamiento de pérdida temporal de cobertura.')
p(d,'Requisitos de red:', bullet=True); p(d,'Conectividad de datos activa mediante Wi Fi corporativa o tarjeta SIM de empresa.', bullet=True); p(d,'DNS, fecha y hora automática configurados correctamente para evitar fallos de certificados y autenticación.', bullet=True); p(d,'Cobertura suficiente en zonas de trabajo o procedimiento operativo de reintento cuando no la haya.', bullet=True); p(d,'No usar redes Wi Fi abiertas para procesos con datos de clientes, firmas o documentos.', bullet=True)

h(d,'5 Seguridad del dispositivo')
table(d,['Control','Exigencia'],[
['Acceso','PIN de al menos seis dígitos, contraseña o biometría; bloqueo automático en un máximo de cinco minutos.'],
['Cifrado','Cifrado del almacenamiento habilitado y actualización de seguridad dentro del plazo definido por sistemas.'],
['Cuentas','Cuenta corporativa individual. Quedan prohibidas cuentas compartidas y credenciales guardadas en notas o aplicaciones no aprobadas.'],
['Aplicaciones','Instalación restringida al catálogo corporativo. No se permite root, depuración USB persistente ni fuentes desconocidas.'],
['Datos','Las fotografías, firmas y documentos deben permanecer en las ubicaciones gestionadas por la aplicación; no se copian a servicios personales.'],
['Incidentes','Pérdida, robo, daño o comportamiento anómalo se comunica de inmediato para bloquear, localizar o borrar el equipo según corresponda.']], [3.1,12.3])

h(d,'6 Operación y soporte')
p(d,'Antes de iniciar la ruta, el usuario debe comprobar batería, cobertura, acceso a la aplicación y permisos de cámara y ubicación. Al finalizar, debe conectar el equipo a carga y comunicar cualquier servicio pendiente de sincronización.')
table(d,['Situación','Actuación requerida'],[
['Sin conexión','Conservar los datos capturados según el comportamiento de la aplicación y reintentar cuando se recupere la red. No duplicar cierres de servicio.'],
['Batería baja','Cargar en vehículo o punto autorizado. No usar cargadores dañados o de procedencia desconocida.'],
['Aplicación no responde','Cerrar y abrir la aplicación; si persiste, reiniciar el dispositivo y contactar con soporte aportando hora, servicio y captura si es posible.'],
['Equipo sustituido','Retirar el dispositivo anterior de la gestión, verificar la sincronización y configurar el reemplazo antes de entregar al usuario.'],
['Actualización','Instalar en una ventana controlada, validar inicio, conexión y una operación básica antes de extenderla al resto de dispositivos.']], [3.3,12.1])

h(d,'7 Validación antes de entrega')
p(d,'Cada dispositivo debe superar la siguiente lista de comprobación antes de asignarse:')
for text in ['Android 7.0 API 24 o superior confirmado.', 'Número de serie, IMEI si aplica, usuario y responsable registrados.', 'Aplicación instalada y autenticación correcta.', 'Conexión con la API validada por Wi Fi y datos móviles.', 'Cámara, ubicación, notificaciones y almacenamiento probados.', 'Bloqueo de pantalla, cifrado y gestión remota activos.', 'Carga de batería, cargador y funda revisados.', 'Prueba de creación y cierre controlado de un servicio, sin usar datos productivos no autorizados.']:
    p(d,text,bullet=True)

h(d,'8 Recomendación de compra')
p(d,'Para nuevas adquisiciones se recomienda seleccionar tabletas Android rugerizadas de 8 pulgadas, 4 GB de RAM, 64 GB de almacenamiento, GPS, 4G o 5G, cámara trasera de 12 MP o superior, batería de 6.000 mAh o más, protección IP65 y soporte del fabricante por al menos tres años. El modelo final debe validarse con una prueba piloto en ruta antes de una compra masiva.')

f=s.footer.paragraphs[0]; f.alignment=WD_ALIGN_PARAGRAPH.CENTER; run_style(f.add_run('ST1 Savall | Requisitos de aplicación hardware y software'),False,8,'808080')
d.save(OUT); print(OUT)
