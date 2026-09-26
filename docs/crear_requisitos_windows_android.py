from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

out=Path(__file__).with_name('requisitos_aplicacion_windows_android_hardware_software.docx')
BLUE='17365D'; GRID='D9D9D9'
def style(r,b=False,s=10,c=None):
 r.font.name='Aptos'; r.font.size=Pt(s); r.bold=b; r._element.rPr.rFonts.set(qn('w:ascii'),'Aptos'); r._element.rPr.rFonts.set(qn('w:hAnsi'),'Aptos')
 if c:r.font.color.rgb=RGBColor.from_string(c)
def fmt(c,fill=None):
 p=c._tc.get_or_add_tcPr()
 if fill:
  x=OxmlElement('w:shd');x.set(qn('w:fill'),fill);p.append(x)
 b=OxmlElement('w:tcBorders')
 for n in ('top','left','bottom','right','insideH','insideV'):
  x=OxmlElement('w:'+n);x.set(qn('w:val'),'single');x.set(qn('w:sz'),'6');x.set(qn('w:color'),GRID);b.append(x)
 p.append(b);m=OxmlElement('w:tcMar')
 for n in ('top','start','bottom','end'):
  x=OxmlElement('w:'+n);x.set(qn('w:w'),'115');x.set(qn('w:type'),'dxa');m.append(x)
 p.append(m);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
def para(d,text,bullet=False):
 p=d.add_paragraph(style='List Bullet' if bullet else None);p.paragraph_format.space_after=Pt(4);p.paragraph_format.line_spacing=1.1;style(p.add_run(text));return p
def head(d,text):
 p=d.add_paragraph(style='Heading 1');p.paragraph_format.space_before=Pt(12);p.paragraph_format.space_after=Pt(5);style(p.add_run(text),True,13.5,'000000')
def tbl(d,hs,rows,w):
 t=d.add_table(rows=1,cols=len(hs));t.style='Table Grid';t.alignment=WD_TABLE_ALIGNMENT.CENTER
 for i,x in enumerate(hs):
  c=t.rows[0].cells[i];c.text='';fmt(c,BLUE);style(c.paragraphs[0].add_run(x),True,9.2,'FFFFFF')
 for ri,row in enumerate(rows):
  cells=t.add_row().cells
  for i,x in enumerate(row):
   c=cells[i];c.text='';fmt(c,'F4F7FA' if ri%2 else None);p=c.paragraphs[0];p.paragraph_format.space_after=Pt(0);style(p.add_run(x),False,9.1)
 for r in t.rows:
  for i,x in enumerate(w):r.cells[i].width=Cm(x)
 d.add_paragraph().paragraph_format.space_after=Pt(2)
d=Document();sec=d.sections[0];sec.top_margin=Cm(2.1);sec.bottom_margin=Cm(2);sec.left_margin=Cm(2.15);sec.right_margin=Cm(2.15)
title=d.add_paragraph(style='Title');title.paragraph_format.space_after=Pt(8);style(title.add_run('Requisitos clave de la aplicación Windows y Android'),True,21,'000000')
sub=d.add_paragraph();sub.paragraph_format.space_after=Pt(16);style(sub.add_run('Hardware y software para ST1 Savall | Versión inicial | 19 de septiembre de 2026'),False,11,'595959')
para(d,'Este documento establece los requisitos para implantar la aplicación ST1 Savall en equipos Windows de oficina y dispositivos Android de campo. Incluye hardware, sistema operativo, software base, red, seguridad, gestión y validación de cada puesto antes de la entrega.')
head(d,'1 Plataformas y alcance')
tbl(d,['Plataforma','Uso previsto','Mínimo de sistema operativo'],[
 ['Windows','Portal de gestión, administración, consultas y operaciones de oficina.','Windows 10 versión 22H2 o Windows 11, 64 bits, con soporte vigente.'],
 ['Android','Ejecución de servicios en campo, captura de datos, firma, fotografías y ubicación.','Android 7.0 o superior, API level 24.']], [3,7,5.3])
head(d,'2 Requisitos de hardware Windows')
tbl(d,['Componente','Mínimo','Recomendado'],[
 ['Procesador','Intel Core i3 o equivalente, 2 núcleos.','Intel Core i5 o equivalente, 4 núcleos.'],
 ['Memoria','8 GB RAM.','16 GB RAM para multitarea y documentos.'],
 ['Almacenamiento','SSD con 20 GB libres.','SSD de 256 GB o superior, con 40 GB libres.'],
 ['Pantalla','13 pulgadas, 1366 x 768.','15 pulgadas, 1920 x 1080 o monitor externo.'],
 ['Red','Ethernet o Wi Fi 5.','Ethernet gigabit o Wi Fi 6; VPN si el acceso es remoto.'],
 ['Periféricos','Teclado, ratón y audio funcional si se usa soporte remoto.','Cámara y auriculares para asistencia y reuniones.']], [3,5.8,6.5])
head(d,'3 Requisitos de hardware Android')
tbl(d,['Componente','Mínimo','Recomendado para campo'],[
 ['Procesador y memoria','4 núcleos y 3 GB RAM.','8 núcleos y 4 GB RAM o superior.'],
 ['Almacenamiento','32 GB internos y 8 GB libres.','64 GB o superior para fotos y actualizaciones.'],
 ['Pantalla','5,5 pulgadas HD.','Tableta de 8 pulgadas Full HD y alto brillo.'],
 ['Cámara y GPS','Cámara trasera con enfoque y GPS funcional.','12 MP, flash y GNSS de buena recepción.'],
 ['Conectividad','Wi Fi y 4G LTE.','4G o 5G, Wi Fi doble banda y Bluetooth 5.'],
 ['Batería y protección','Carga para una jornada y funda.','6.000 mAh o más, cargador de vehículo e IP65 o superior.']], [3,5.8,6.5])
head(d,'4 Software y configuración Windows')
tbl(d,['Área','Requisito'],[
 ['Navegador','Microsoft Edge o Google Chrome en versión soportada y actualizada.'],
 ['Acceso web','URL productiva de la aplicación accesible por HTTPS desde la red autorizada.'],
 ['Documentos','Visor PDF actualizado; suite ofimática corporativa si se exportan o revisan ficheros.'],
 ['Seguridad','Microsoft Defender o protección corporativa activa, cifrado de disco y bloqueo automático de sesión.'],
 ['Gestión','Equipo unido al dominio o MDM, inventariado y con actualizaciones de Windows gestionadas.'],
 ['Distribución','Instalar solo software y extensiones autorizadas; no usar cuentas locales compartidas.']], [3.3,12])
head(d,'5 Software y configuración Android')
tbl(d,['Área','Requisito'],[
 ['Aplicación','Versión aprobada de ST1 Savall instalada desde el canal corporativo y asociada al usuario.'],
 ['Permisos','Cámara, ubicación, almacenamiento y notificaciones habilitados solo para la función necesaria.'],
 ['Servicios base','Fecha y hora automáticas, Google Play Services cuando la distribución lo requiera y navegador actualizado.'],
 ['Gestión remota','Bloqueo, cifrado, inventario, actualización y borrado remoto mediante MDM cuando esté disponible.'],
 ['Restricciones','Sin root, ROM modificada, depuración USB persistente ni instalación desde fuentes desconocidas.']], [3.3,12])
head(d,'6 Red y seguridad comunes')
para(d,'Los equipos Windows y Android deben acceder a la API corporativa exclusivamente mediante HTTPS. La configuración se separa por entorno y no incluye contraseñas, tokens ni cadenas de conexión en la aplicación o en archivos compartidos.')
for x in ['Conexión por red corporativa, VPN autorizada o datos móviles de empresa.', 'Cuentas individuales y permisos por función; prohibidas las credenciales compartidas.', 'Bloqueo automático: máximo cinco minutos en Android y quince minutos en Windows.', 'Cifrado de almacenamiento y actualización de seguridad dentro del plazo definido por sistemas.', 'Prohibido copiar firmas, fotografías o documentos de trabajo a correos, nubes o dispositivos personales.', 'Comunicar pérdida, robo, daño o comportamiento anómalo de inmediato para su bloqueo y análisis.']:para(d,x,True)
head(d,'7 Operación y soporte')
tbl(d,['Situación','Actuación'],[
 ['Actualización de aplicación','Publicar primero en piloto, validar inicio, autenticación y una operación básica; después extender al resto.'],
 ['Sin conexión','Reintentar cuando vuelva la red. No repetir operaciones de cierre sin comprobar el resultado previo.'],
 ['Incidencia de acceso','Registrar usuario, equipo, hora, captura si procede y mensaje de error; soporte revisa permisos, red y versión.'],
 ['Sustitución de equipo','Verificar sincronización, retirar el equipo previo de la gestión y validar el nuevo antes de entregarlo.'],
 ['Mantenimiento','Revisar semanalmente actualizaciones, capacidad libre, errores repetidos y estado de protección.']], [3.5,11.8])
head(d,'8 Lista de validación antes de entrega')
for x in ['Equipo inventariado y asignado a un usuario y responsable.', 'Sistema operativo mínimo confirmado: Windows 10 22H2 o Windows 11, o Android 7 API 24.', 'Aplicación o acceso web instalado y autenticación correcta.', 'Conexión validada desde la red habitual de trabajo.', 'Cámara, ubicación y notificaciones comprobadas en Android.', 'Navegador, visor PDF, protección antimalware y actualizaciones comprobadas en Windows.', 'Bloqueo, cifrado y gestión remota activos.', 'Prueba funcional controlada completada y resultado registrado.']:para(d,x,True)
foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.CENTER;style(foot.add_run('ST1 Savall | Requisitos Windows y Android'),False,8,'808080')
d.save(out);print(out)
