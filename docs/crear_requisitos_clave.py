from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path(__file__).with_name('requisitos_clave_implementacion_operatividad.docx')

BLUE = '17365D'
LIGHT_BLUE = 'DCE6F1'
GRID = 'D9D9D9'

def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tc_pr.append(shd)

def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in('w:tcMar')
    if tc_mar is None:
        tc_mar = OxmlElement('w:tcMar')
        tc_pr.append(tc_mar)
    for side, value in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tc_mar.find(qn(f'w:{side}'))
        if node is None:
            node = OxmlElement(f'w:{side}')
            tc_mar.append(node)
        node.set(qn('w:w'), str(value))
        node.set(qn('w:type'), 'dxa')

def set_cell_border(cell, color=GRID):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        tc_pr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        tag = qn(f'w:{edge}')
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f'w:{edge}')
            borders.append(element)
        element.set(qn('w:val'), 'single')
        element.set(qn('w:sz'), '6')
        element.set(qn('w:color'), color)

def style_run(run, bold=False, size=10.5, color=None):
    run.font.name = 'Aptos'
    run._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos')
    run._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)

def add_body(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.12
    style_run(p.add_run(text))
    return p

def add_bullet(doc, text):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(3)
    style_run(p.add_run(text))
    return p

def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f'Heading {level}')
    p.paragraph_format.space_before = Pt(13 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    style_run(r, bold=True, size=14 if level == 1 else 11.5, color='000000')
    return p

def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'
    hdr = table.rows[0]
    for i, text in enumerate(headers):
        cell = hdr.cells[i]
        cell.text = ''
        set_cell_shading(cell, BLUE)
        set_cell_margins(cell)
        set_cell_border(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        style_run(p.add_run(text), bold=True, size=9.5, color='FFFFFF')
    for index, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cell = cells[i]
            cell.text = ''
            if index % 2 == 1:
                set_cell_shading(cell, 'F4F7FA')
            set_cell_margins(cell)
            set_cell_border(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            style_run(p.add_run(value), size=9.3)
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Cm(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table

doc = Document()
section = doc.sections[0]
section.top_margin = Cm(2.1)
section.bottom_margin = Cm(2.0)
section.left_margin = Cm(2.15)
section.right_margin = Cm(2.15)

styles = doc.styles
styles['Normal'].font.name = 'Aptos'
styles['Normal']._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos')
styles['Normal']._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
styles['Normal'].font.size = Pt(10.5)

title = doc.add_paragraph(style='Title')
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
title.paragraph_format.space_after = Pt(8)
style_run(title.add_run('Requisitos clave para implementación y operatividad'), bold=True, size=22, color='000000')
subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(18)
style_run(subtitle.add_run('Sistema ST1 Savall | Versión inicial | 19 de septiembre de 2026'), size=11, color='595959')

add_body(doc, 'Este documento fija los requisitos mínimos para implantar, poner en marcha y operar de forma segura la solución ST1 Savall. Su propósito es alinear a desarrollo, sistemas y responsables operativos antes de cualquier despliegue productivo. Cubre la aplicación móvil, el portal web, la API, las bases de datos SQL Server y las integraciones con Sage 50; la integración de telemetría Wialon queda condicionada a su contratación y configuración.')

add_heading(doc, '1 Alcance y arquitectura')
add_body(doc, 'La solución está formada por una aplicación móvil .NET MAUI para personal de campo, un portal web Blazor para gestión, una API ASP.NET Core y un componente compartido de modelos e interfaz. La persistencia se apoya en SQL Server y la API consulta datos operativos propios y datos de Sage 50. Las operaciones críticas incluyen la gestión de servicios, obras, artículos, precios especiales, albaranes y documentación asociada.')
add_table(doc, ['Componente', 'Responsabilidad', 'Requisito de disponibilidad'], [
    ['Aplicación móvil', 'Ejecución de tareas de campo, captura de datos y cierre de servicios.', 'Debe funcionar en dispositivos homologados y gestionar errores de conectividad.'],
    ['Portal web', 'Administración, consulta y operaciones de oficina.', 'Acceso mediante navegador compatible y red corporativa o acceso remoto protegido.'],
    ['API', 'Reglas de negocio, autenticación, integración con SQL Server, Sage 50 y generación documental.', 'Servicio central disponible antes de habilitar clientes.'],
    ['SQL Server', 'Datos de la aplicación y bases de datos de Sage 50.', 'Copias de seguridad, acceso restringido y capacidad suficiente.'],
    ['Servicios externos', 'Correo, generación de PDFs y, si aplica, Wialon.', 'Configuración separada por entorno y monitorización de fallos.']
], [3.1, 7.3, 5.0])

add_heading(doc, '2 Requisitos previos de implementación')
add_table(doc, ['Área', 'Requisito mínimo', 'Criterio de aceptación'], [
    ['Infraestructura', 'Servidor Windows compatible con .NET 9, SQL Server accesible y almacenamiento para adjuntos, exportaciones y registros.', 'La API arranca en el entorno objetivo y se conecta a sus dependencias sin usar valores de desarrollo.'],
    ['Red', 'Resolución DNS, puertos permitidos entre clientes, API, SQL Server, Sage 50 y SMTP. Acceso cifrado cuando el tráfico sale de la red de confianza.', 'Pruebas de conectividad desde cada rol de usuario y desde el servidor de la API.'],
    ['Identidad', 'Cuentas individuales, grupos por función y autenticación integrada o mecanismo equivalente aprobado.', 'Un usuario sin permiso no puede acceder a pantallas ni endpoints restringidos.'],
    ['Datos', 'Inventario de empresas, obras, clientes, artículos, series, almacenes y parámetros necesarios para los procesos de negocio.', 'Muestra de datos validada por negocio antes de la carga completa.'],
    ['Dispositivos', 'Tabletas o teléfonos compatibles, cámara operativa, hora automática y conectividad estable o procedimiento offline definido.', 'Prueba de un servicio completo desde un dispositivo representativo.'],
    ['Licencias', 'Licencias vigentes de DevExpress, SQL Server, Sage 50 y cualquier servicio de mapas, correo o telemetría usado.', 'Titularidad, vencimiento y responsable registrados.']
], [2.5, 7.4, 5.5])

add_heading(doc, '3 Configuración por entornos')
add_body(doc, 'Deben existir, como mínimo, los entornos de desarrollo, pruebas y producción. Producción no debe reutilizar bases de datos, credenciales, buzones de correo ni claves de los demás entornos. Toda configuración debe residir fuera del código fuente, mediante secretos gestionados o variables protegidas.')
add_bullet(doc, 'Registrar las cadenas de conexión, credenciales SMTP, tokens de terceros y claves de firma como secretos; no incluirlos en repositorios, documentos ni capturas de pantalla.')
add_bullet(doc, 'Mantener una ficha de configuración por entorno con propietario, fecha de revisión, dependencias, URL pública o interna y procedimiento de renovación de secretos.')
add_bullet(doc, 'Configurar las URL base de la API y las políticas CORS de forma explícita para los clientes autorizados.')
add_bullet(doc, 'Usar configuración de registro diferenciada: suficiente diagnóstico en pruebas y niveles controlados en producción, sin datos personales o secretos innecesarios.')

add_heading(doc, '4 Integración con Sage 50')
add_body(doc, 'La API es el único componente autorizado para comunicarse con Sage 50. La aplicación móvil y el portal web no deben conectarse directamente a las bases de datos de Sage. Los permisos técnicos han de limitarse a las tablas, operaciones y empresas estrictamente necesarias.')
add_table(doc, ['Control', 'Requisito operativo'], [
    ['Conectividad', 'Validar acceso a las bases de gestión y comunes de Sage desde el servidor de la API antes del despliegue.'],
    ['Datos maestros', 'Comprobar la correspondencia entre obras, clientes, artículos, series, almacenes e impuestos antes de activar operaciones transaccionales.'],
    ['Albaranes', 'La emisión debe ser transaccional: reservar numeración, crear cabecera y líneas, y registrar la referencia en el servicio. Si falla, el servicio no se da por finalizado.'],
    ['Idempotencia', 'Un reintento no puede generar un segundo albarán. La serie y número generados deben quedar vinculados al servicio.'],
    ['Cambios de esquema', 'Cualquier actualización de Sage debe probarse primero en entorno no productivo y validar las consultas de integración.'],
    ['Recuperación', 'Definir responsable y procedimiento para incidencias de conectividad, bloqueo de series, datos maestros ausentes o documentos parcialmente emitidos.']
], [3.4, 12.0])

add_heading(doc, '5 Seguridad y protección de datos')
add_bullet(doc, 'Aplicar mínimo privilegio a usuarios, cuentas de servicio y accesos a bases de datos. Las cuentas compartidas solo se permitirán con justificación, propietario y revisión periódica.')
add_bullet(doc, 'Exigir contraseñas robustas o autenticación corporativa, bloquear el acceso cuando se revoca una cuenta y conservar trazabilidad de altas, bajas y cambios de permisos.')
add_bullet(doc, 'Proteger el tráfico externo con HTTPS y certificados renovados antes de su vencimiento. Restringir el acceso administrativo a redes y equipos autorizados.')
add_bullet(doc, 'No registrar contraseñas, cadenas de conexión, tokens, firmas, documentos de identidad ni adjuntos en texto plano. Enmascarar datos sensibles en diagnósticos.')
add_bullet(doc, 'Definir plazos de conservación para servicios, albaranes, firmas, fotografías y registros, junto con el proceso aprobado para consulta, exportación y eliminación.')

add_heading(doc, '6 Despliegue y puesta en marcha')
add_table(doc, ['Fase', 'Actividades obligatorias', 'Salida esperada'], [
    ['Preparación', 'Inventario de versiones, respaldo verificable, ventana de cambio, responsables y plan de reversión.', 'Cambio autorizado y recuperable.'],
    ['Publicación', 'Compilar versión Release, desplegar API y clientes, aplicar migraciones de forma controlada y verificar configuración.', 'Servicios iniciados con la versión prevista.'],
    ['Validación técnica', 'Prueba de salud, autenticación, lectura y escritura controlada, generación de PDF, envío de correo si aplica y acceso a Sage.', 'Evidencia de pruebas sin errores bloqueantes.'],
    ['Validación funcional', 'Ejecutar casos de servicio, cierre, emisión de albarán, consulta en portal y verificación de la trazabilidad.', 'Aprobación de negocio para activar.'],
    ['Activación', 'Piloto con usuarios y flota limitados; ampliar de forma gradual tras comprobar indicadores.', 'Operación productiva estable.'],
    ['Reversión', 'Volver a la versión anterior o deshabilitar la funcionalidad afectada sin perder datos ya confirmados.', 'Servicio restablecido y causa registrada.']
], [2.6, 8.1, 4.7])

add_heading(doc, '7 Requisitos de operación diaria')
add_table(doc, ['Frecuencia', 'Control', 'Responsable'], [
    ['Diaria', 'Revisar disponibilidad de API, errores de integración, tareas de servicios pendientes y emisión de albaranes.', 'Operación / soporte de primer nivel'],
    ['Diaria', 'Comprobar que los cierres de servicio tienen datos obligatorios, firma y referencia de albarán cuando corresponde.', 'Responsable operativo'],
    ['Semanal', 'Revisar fallos de correo, registros de error, espacio disponible, cuentas bloqueadas y calidad de sincronización.', 'Sistemas'],
    ['Mensual', 'Verificar restauración de copia de seguridad, aplicar parches planificados y revisar permisos y secretos próximos a vencer.', 'Sistemas y seguridad'],
    ['Por cambio', 'Actualizar inventario, documentación, pruebas y plan de reversión antes de modificar procesos, bases de datos o integraciones.', 'Propietario del cambio']
], [2.2, 9.1, 4.1])

add_heading(doc, '8 Monitorización, copias y continuidad')
add_body(doc, 'La monitorización debe detectar indisponibilidad de la API, errores de autenticación, fallos de SQL Server, interrupciones con Sage, saturación de disco, errores de envío y generación documental. Las alertas han de incluir fecha, componente, gravedad, correlación del error y canal de escalado, sin revelar información confidencial.')
add_bullet(doc, 'Realizar copias de seguridad periódicas de la base de datos de la aplicación, ficheros documentales y configuraciones críticas, conforme a la política corporativa.')
add_bullet(doc, 'Probar la restauración en un entorno aislado con una periodicidad definida; una copia no se considera válida hasta haber sido restaurada y verificada.')
add_bullet(doc, 'Documentar RPO y RTO aprobados por negocio, así como el orden de recuperación: datos, API, integraciones, clientes y validación funcional.')

add_heading(doc, '9 Criterios de aceptación para producción')
add_table(doc, ['Ámbito', 'Condición de aceptación'], [
    ['Funcional', 'Los casos de alta, asignación, ejecución y cierre de servicio se completan con datos consistentes.'],
    ['Sage 50', 'La emisión de albarán crea un único documento correcto, actualiza la numeración y deja la referencia trazable en el servicio.'],
    ['Seguridad', 'Los secretos están fuera del código, los permisos se han probado por rol y los accesos no autorizados son rechazados.'],
    ['Rendimiento', 'Las operaciones habituales responden dentro de los niveles acordados y no se observan bloqueos recurrentes en base de datos.'],
    ['Resiliencia', 'Se ha probado la recuperación ante caída de API, pérdida de conexión con Sage y reintento de una operación crítica.'],
    ['Operación', 'Existe responsable, guía de incidencias, procedimiento de copias y contactos de escalado disponibles para el equipo.']
], [3.1, 12.3])

add_heading(doc, '10 Pendientes de decisión antes de ampliar alcance')
add_bullet(doc, 'Confirmar los valores que Sage debe heredar del cliente o artículo en los albaranes: forma de pago, IVA, cuenta, familia, vendedor y otras reglas específicas.')
add_bullet(doc, 'Si se activa Wialon, confirmar modalidad, URL de acceso, token técnico de solo lectura, asociación única camión-unidad y disponibilidad de Wialon Logistics para rutas planificadas.')
add_bullet(doc, 'Acordar objetivos de disponibilidad, tiempo máximo de recuperación, retención de registros y requisitos de trabajo sin cobertura móvil.')
add_bullet(doc, 'Designar propietarios de negocio, sistemas, seguridad y soporte de primer nivel, con sustitutos y horario de escalado.')

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
style_run(footer.add_run('ST1 Savall | Requisitos clave de implementación y operatividad'), size=8, color='808080')

doc.save(OUT)
print(OUT)
